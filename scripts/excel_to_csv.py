"""Chuyển tab PRODUCT_MASTER trong file Excel kế hoạch → data/products.csv cho bot.

Chạy: python scripts/excel_to_csv.py "Plan Shopee An Lạc Phát Lighting.xlsx"
Chỉ xuất thông tin được phép cho khách xem: KHÔNG xuất giá nhập, lợi nhuận, margin, số tồn kho.
"""
import csv, re, sys, pathlib
import openpyxl

NOT_ON_PLATFORM = {"Opple"}  # thương hiệu chưa lên sàn
# Sửa lỗi chính tả tên sản phẩm trong file gốc (khớp theo tên gốc)
NAME_FIX = {
    "x2 đôi công tắc 3 chấu + mặt nạ": "2 đôi ổ cắm 3 chấu + 1 mặt nạ",
    "Đôi ổ cắm 3 chấu + 3 công tắt + mặt nạ": "Đôi ổ cắm 3 chấu + 3 công tắc + 1 mặt nạ",
    "Đôi ỏ cắm 3 chấu + 1 mặt nạ": "Đôi ổ cắm 3 chấu + 1 mặt nạ",
}
OUT = pathlib.Path(__file__).resolve().parent.parent / "data" / "products.csv"


def round_k(v):
    return f"{int(round(float(v) / 1000.0)) * 1000:,}".replace(",", ".") + "đ" if v else ""


def clean_note(note):
    keep = []
    for part in str(note or "").split(";"):
        p = part.strip()
        if not p or p.startswith(("⚠", "🟢", "🔴", "🟡")):
            continue
        p = re.sub(r"^Màu:\s*(Màu:\s*|—Màu:\s*)?", "Màu: ", p)
        keep.append(p)
    return "; ".join(dict.fromkeys(keep))


def main(path):
    ws = openpyxl.load_workbook(path, data_only=True)["PRODUCT_MASTER"]
    header_row = next(i for i, r in enumerate(ws.iter_rows(values_only=True), 1) if r[0] == "STT")
    idx = {h: i for i, h in enumerate(next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))) if h}
    rows = []
    for r in ws.iter_rows(min_row=header_row + 1, values_only=True):
        sku, name = r[idx["SKU"]], r[idx["Sản phẩm"]]
        if not sku or not name:
            continue
        note = str(r[idx["Ghi chú"]] or "")
        if "Dòng nháp" in note:  # dòng chưa có giá chính thức
            continue
        brand = r[idx["Thương hiệu"]]
        stock = r[idx["Tồn kho"]]
        status = r[idx["Cảnh báo"]] or ""
        het = "Hết hàng" in status or (stock is not None and stock <= 0)
        rows.append({
            "ma_sp": "" if "chưa có mã" in str(sku) else re.sub(r"\s*\+\s*", " + ", str(sku).strip()),
            "ten_san_pham": NAME_FIX.get(str(name).strip(), str(name).strip()),
            "thuong_hieu": brand,
            "loai": r[idx["Nhóm hàng"]],
            "gia_shopee": round_k(r[idx["Giá bán Shopee (đ)"]]),
            "gia_tiktok": round_k(r[idx["Giá bán TikTok (đ)"]]),
            "tinh_trang": "Đặt trước (kho tạm hết, nhập theo yêu cầu)" if het else "Có sẵn",
            "ban_tren_san": "Chưa (mua trực tiếp)" if brand in NOT_ON_PLATFORM else "Có",
            "ghi_chu": clean_note(note),
        })
    with open(OUT, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)
    print(f"Đã xuất {len(rows)} sản phẩm → {OUT}")


if __name__ == "__main__":
    main(sys.argv[1])
